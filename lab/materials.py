"""Local document snapshots. Import never modifies the original or runs macros."""
import hashlib
import io
from pathlib import Path
from zipfile import ZipFile, BadZipFile
import xml.etree.ElementTree as ET
from core import now

MAX_FILE = 10 * 1024 * 1024
MAX_TEXT = 120000
MAX_CONTEXT_BYTES = 18000


def extract(path):
    path = Path(path)
    if path.suffix.lower() not in ('.txt', '.md', '.csv', '.docx', '.xlsx'):
        raise ValueError('Admite TXT, Markdown, CSV, DOCX y XLSX. Convierte los formatos antiguos primero.')
    with path.open('rb') as stream:
        data = stream.read(MAX_FILE + 1)
    if len(data) > MAX_FILE:
        raise ValueError('El archivo supera 10 MB.')
    suffix = path.suffix.lower()
    warnings = []
    if suffix in ('.docx', '.xlsx'):
        try:
            with ZipFile(io.BytesIO(data)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 40 * 1024 * 1024 or len(archive.infolist()) > 3000:
                    raise ValueError('Documento demasiado grande al descomprimir.')
                if suffix == '.docx':
                    xml = archive.read('word/document.xml')
                    if b'<!DOCTYPE' in xml or b'<!ENTITY' in xml:
                        raise ValueError('XML no admitido.')
                    root = ET.fromstring(xml)
                    ns = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
                    lines = [''.join(t.text or '' for t in p.iter(ns+'t')) for p in root.iter(ns+'p')]
                    text = '\n'.join(lines)
                    warnings.append('Solo texto del cuerpo y tablas; sin imágenes, encabezados, notas ni formato visual.')
                else:
                    from openpyxl import load_workbook
                    book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
                    lines = []; size = 0; count = 0
                    try:
                        for sheet in book:
                            if (sheet.max_row or 0)*(sheet.max_column or 0)>100000:
                                raise ValueError('Excel supera el límite de 100000 celdas por hoja.')
                            lines.append(f'Hoja: {sheet.title}')
                            for row in sheet.iter_rows():
                                count += len(row)
                                if count > 100000:raise ValueError('Excel supera el límite de 100000 celdas leídas.')
                                if not any(c.value is not None for c in row):continue
                                line = ' | '.join(f'{c.coordinate}: {c.value}' for c in row if c.value is not None)
                                size += len(line)
                                if size > MAX_TEXT:raise ValueError('Demasiado texto; importa un documento más pequeño.')
                                lines.append(line)
                    finally:book.close()
                    text = '\n'.join(lines)
                    warnings.append('Valores y fórmulas como texto; no se calculan fórmulas ni se ejecutan macros. Sin imágenes ni formato.')
        except (BadZipFile, KeyError, ET.ParseError) as exc:
            raise ValueError('Documento inválido o dañado.') from exc
    else:
        try:text = data.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = data.decode('cp1252');warnings.append('Texto leído con codificación Windows-1252.')
        if '\x00' in text:raise ValueError('El archivo no parece texto.')
        if suffix == '.csv':
            warnings.append('CSV conservado como texto; no modifica la lista de alumnos.')
    text = text.strip()
    if not text:raise ValueError('No se encontró texto legible.')
    if len(text) > MAX_TEXT:raise ValueError('Demasiado texto; importa un documento más pequeño.')
    return text, '\n'.join(warnings), hashlib.sha256(data).hexdigest()


class Materials:
    def __init__(self, store):
        self.store = store
        store.con.executescript('''
        CREATE TABLE IF NOT EXISTS materials(
          id INTEGER PRIMARY KEY, group_id INTEGER REFERENCES groups(id), title TEXT NOT NULL,
          content TEXT NOT NULL, warning TEXT NOT NULL, digest TEXT NOT NULL, created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS message_materials(
          message_id INTEGER NOT NULL REFERENCES messages(id), material_id INTEGER NOT NULL REFERENCES materials(id),
          PRIMARY KEY(message_id,material_id));
        ''')

    def save(self, gid, title, content, warning='Material guardado por el docente; revisar antes de usar.', digest=None):
        title = title.strip();content = content.strip()
        if not title or not content:raise ValueError('Escribe título y contenido.')
        if len(title)>200 or len(content)>MAX_TEXT:raise ValueError('Material demasiado largo.')
        digest = digest or hashlib.sha256(content.encode()).hexdigest()
        with self.store.con:
            row = self.store.con.execute('SELECT id FROM materials WHERE group_id IS ? AND digest=?',(gid,digest)).fetchone()
            if row:return row[0]
            return self.store.con.execute('INSERT INTO materials(group_id,title,content,warning,digest,created_at) VALUES(?,?,?,?,?,?)',
                                          (gid,title,content,warning,digest,now())).lastrowid

    def import_file(self,gid,path):
        text,warning,digest = extract(path)
        return self.save(gid,Path(path).name,text,warning,digest)

    def list(self,gid,query=''):
        rows=self.store.con.execute('SELECT id,title,created_at FROM materials WHERE group_id IS ? ORDER BY id DESC',(gid,))
        return [dict(r) for r in rows if query.casefold() in r['title'].casefold()]

    def get(self,gid,mid):
        row=self.store.con.execute('SELECT * FROM materials WHERE id=? AND group_id IS ?',(mid,gid)).fetchone()
        if row is None:raise ValueError('Material no disponible en este grupo.')
        return dict(row)

    def context(self,gid,ids):
        ids=list(dict.fromkeys(ids))
        if len(ids)>3:raise ValueError('Adjunta como máximo tres materiales por consulta.')
        result=[]
        for mid in ids:
            r=self.get(gid,mid)
            result.append({'source':f'M{mid}','title':r['title'],'text':r['content'],'notice':r['warning']})
        import json
        if len(json.dumps(result,ensure_ascii=False).encode())>MAX_CONTEXT_BYTES:
            raise ValueError('Los adjuntos superan 18 KB de contexto. Guarda un fragmento más pequeño como material; no se enviaron recortes ocultos.')
        return result
