"""Synthetic, offline packaging smoke checks. No user DB, credentials or microphone."""
import json
import os
import platform
import tempfile
from pathlib import Path
from zipfile import ZipFile
from version import VERSION


def run_checks(check_gui=False):
    checks = []

    def check(name, operation):
        try:
            operation()
            checks.append({'name': name, 'status': 'passed'})
        except Exception as exc:
            # Keep private paths, payloads and credentials out of portable reports.
            checks.append({'name': name, 'status': 'failed', 'error_type': type(exc).__name__})

    with tempfile.TemporaryDirectory(prefix='aulaos-check-') as folder:
        root = Path(folder)

        def storage():
            from core import Store
            db = Store(root / 'synthetic.sqlite3')
            try:
                gid = db.group('Escuela sintética', 'Prueba')
                sid = db.add_student(gid, 'Estudiante ficticio')
                db.mark(gid, sid, '2026-01-01', 'A')
                assert db.marks(gid, '2026-01-01')[sid] == 'A'
                backup = db.backup()
            finally:
                db.close()
            copy = Store(backup)
            try:
                assert copy.marks(gid, '2026-01-01')[sid] == 'A'
            finally:
                copy.close()
        check('sqlite_backup_reopen', storage)

        def documents():
            from docx import Document
            from templates import Templates
            from core import Store
            original = root / 'format.docx'
            doc = Document()
            doc.add_table(rows=2, cols=3).cell(0, 0).text = 'Actividad'
            doc.save(original)
            db = Store(root / 'documents.sqlite3')
            try:
                gid = db.group('Escuela sintética', 'Formato')
                templates = Templates(db)
                tid = templates.import_file(gid, original, 'Tres columnas')
                did = templates.create(gid, tid, 'Clase')
                templates.export(gid, did, root / 'export.docx')
                assert len(Document(root / 'export.docx').tables[0].columns) == 3
            finally:
                db.close()
        check('docx_template_export', documents)

        def spreadsheet():
            from openpyxl import Workbook, load_workbook
            path = root / 'sheet.xlsx'
            book = Workbook()
            book.active.append(['Nombre', 'Matrícula'])
            book.active.append(['Estudiante ficticio', '0001'])
            book.save(path)
            book.close()
            book = load_workbook(path)
            try:
                assert book.active['B2'].value == '0001'
            finally:
                book.close()
        check('xlsx_read_write', spreadsheet)

        def slides():
            from lesson_export import export_presentation
            data = dict(title='Clase sintética', teacher_guide='Guía docente', worksheet='Actividad', slides=[
                dict(title='Inicio', body=['Read and speak.'], notes='Presenta la actividad.'),
                dict(title='Cierre', body=['Share one sentence.'], notes='Revisa la participación.')])
            path = root / 'slides.pptx'
            export_presentation(data, 'Comprobación sintética', path)
            with ZipFile(path) as archive:
                assert archive.testzip() is None
                assert 'ppt/slides/slide2.xml' in archive.namelist()
        check('pptx_bundled_asset', slides)

        if os.name == 'nt':
            def credentials():
                from settings import _dpapi
                assert _dpapi(_dpapi(b'synthetic-check', True), False) == b'synthetic-check'
            check('windows_dpapi', credentials)
            def audio_library():
                import sounddevice
                assert sounddevice.get_portaudio_version()[0] > 0
            check('portaudio_load_only', audio_library)
        else:
            checks.append({'name': 'windows_dpapi_and_portaudio', 'status': 'skipped'})

        if check_gui:
            def gui():
                import tkinter as tk
                window = tk.Tk()
                try:
                    window.withdraw()
                    window.update()
                finally:
                    window.destroy()
            check('tk_initialization_only', gui)
    return dict(version=VERSION, platform=platform.system(), checks=checks,
                ok=all(c['status'] != 'failed' for c in checks),
                limitations=['No real API calls', 'No microphone capture', 'No teacher acceptance test',
                             'No installer/update/uninstall validation'])
