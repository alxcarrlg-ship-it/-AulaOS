"""Export a saved draft to a new Word, printable HTML, or Markdown document.
A deliberately small Markdown subset; unsupported syntax stays as visible text.
"""
import html
import io
import os
from pathlib import Path
import re
import tempfile
from drafts import Drafts
from research import safe_url

INLINE=re.compile(r'\[([^\]\n]+)\]\((https?://[^\s)]+)\)|\*\*([^*\n]+)\*\*')


def inline(text):
    offset=0
    for match in INLINE.finditer(text):
        if match.start()>offset:yield 'text',text[offset:match.start()],None
        if match.group(1) is not None:
            url=safe_url(match.group(2))
            yield ('link',match.group(1),url) if url else ('text',match.group(0),None)
        else:yield 'bold',match.group(3),None
        offset=match.end()
    if offset<len(text):yield 'text',text[offset:],None


def blocks(content):
    lines=content.splitlines();index=0
    def cells(line):return [c.strip() for c in line.strip().strip('|').split('|')]
    while index<len(lines):
        line=lines[index];index+=1
        if not line.strip():continue
        if line.strip().startswith('```'):
            code=[]
            while index<len(lines) and not lines[index].strip().startswith('```'):
                code.append(lines[index]);index+=1
            if index<len(lines):index+=1
            yield 'code','\n'.join(code);continue
        if '|' in line and index<len(lines):
            separator=cells(lines[index]);head=cells(line)
            if len(head)==len(separator) and len(head)>1 and all(re.fullmatch(r':?-{3,}:?',s) for s in separator):
                rows=[head];index+=1
                while index<len(lines) and '|' in lines[index]:
                    row=cells(lines[index])
                    if len(row)!=len(head):break
                    rows.append(row);index+=1
                # Wide tables remain text, preserving every cell instead of clipping.
                if len(head)<=6:yield 'table',rows
                else:
                    for row in rows:yield 'paragraph',' | '.join(row)
                continue
        heading=re.match(r'^(#{1,3})\s+(.+)$',line)
        bullet=re.match(r'^\s*[-*+]\s+(.+)$',line)
        ordered=re.match(r'^\s*(\d+[.)])\s+(.+)$',line)
        if heading:yield 'h'+str(len(heading[1])),heading[2]
        elif bullet:yield 'bullet',bullet[1]
        elif ordered:yield 'number',ordered[1]+' '+ordered[2]
        else:yield 'paragraph',line


def html_inline(text):
    parts=[]
    for kind,value,url in inline(text):
        value=html.escape(value)
        if kind=='link':parts.append(f'<a href="{html.escape(url,quote=True)}" rel="noopener noreferrer">{value}</a>')
        elif kind=='bold':parts.append('<strong>'+value+'</strong>')
        else:parts.append(value)
    return ''.join(parts)


def html_document(title,content):
    body=[]
    for kind,value in blocks(content):
        if kind=='table':
            rows=[]
            for n,row in enumerate(value):
                tag='th' if n==0 else 'td'
                rows.append('<tr>'+''.join(f'<{tag}>'+html_inline(cell)+f'</{tag}>' for cell in row)+'</tr>')
            body.append('<table><thead>'+rows[0]+'</thead><tbody>'+''.join(rows[1:])+'</tbody></table>')
        elif kind=='code':body.append('<pre>'+html.escape(value)+'</pre>')
        elif kind=='bullet':body.append('<p class="bullet">• '+html_inline(value)+'</p>')
        else:
            tag=kind if kind.startswith('h') else 'p'
            body.append(f'<{tag}>'+html_inline(value)+f'</{tag}>')
    return ('''<!doctype html><html lang="es"><head><meta charset="utf-8">
<meta http-equiv="Content-Security-Policy" content="default-src 'none'; style-src 'unsafe-inline'; base-uri 'none'; form-action 'none'">
<title>'''+html.escape(title)+'''</title><style>
@page{size:letter;margin:20mm}*{box-sizing:border-box}body{max-width:175mm;margin:24px auto;font:11pt/1.4 Arial,sans-serif;color:#1e293b}
h1,h2,h3{color:#000000;break-after:avoid}h1{font-size:19pt}h2{font-size:15pt}h3{font-size:12pt}
header{margin-bottom:22px;padding-bottom:10px}header h1{font-size:25pt;margin:0}
p{margin:0 0 9px;white-space:pre-wrap;overflow-wrap:anywhere}.bullet{padding-left:12px}
table{width:100%;border-collapse:collapse;table-layout:fixed;margin:12px 0}th,td{border:1px solid #b8c7d1;padding:7px;vertical-align:top;overflow-wrap:anywhere}
th{background:#e9f2f5;text-align:left}thead{display:table-header-group}pre{white-space:pre-wrap;overflow-wrap:anywhere;border-left:3px solid #176b78;padding:10px}
a{color:#155e75;overflow-wrap:anywhere}@media print{body{margin:0;max-width:none}a{color:inherit}}
</style></head><body><header><h1>'''+html.escape(title)+'</h1></header>'+''.join(body)+'</body></html>').encode('utf-8')


def word_document(title,content):
    from docx import Document
    from docx.shared import Inches,Pt,RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE
    doc=Document();section=doc.sections[0]
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Inches(.8)
    normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.15
    for name,size in [('Title',25),('Heading 1',19),('Heading 2',15),('Heading 3',12)]:
        style=doc.styles[name];style.font.name='Arial';style.font.size=Pt(size);style.font.color.rgb=RGBColor.from_string('000000')
        for border in style.element.xpath('./w:pPr/w:pBdr'):border.getparent().remove(border)
    doc.core_properties.title=title;doc.core_properties.author='';doc.core_properties.last_modified_by=''
    doc.core_properties.comments=''
    doc.add_heading(title,0)
    def add_text(paragraph,text):
        for kind,value,url in inline(text):
            if kind=='link':
                link=OxmlElement('w:hyperlink');link.set(qn('r:id'),paragraph.part.relate_to(url,RELATIONSHIP_TYPE.HYPERLINK,is_external=True))
                run=OxmlElement('w:r');props=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'155E75');props.append(color)
                underline=OxmlElement('w:u');underline.set(qn('w:val'),'single');props.append(underline);run.append(props)
                node=OxmlElement('w:t');node.set(qn('xml:space'),'preserve');node.text=value;run.append(node);link.append(run);paragraph._p.append(link)
            else:
                run=paragraph.add_run(value);run.bold=kind=='bold'
    for kind,value in blocks(content):
        if kind=='table':
            table=doc.add_table(rows=1,cols=len(value[0]));table.style='Table Grid';table.autofit=False
            weights=[max(8,min(50,max(len(row[i]) for row in value))) for i in range(len(value[0]))]
            widths=[Inches(6.9*(.55/len(weights)+.45*w/sum(weights))) for w in weights]
            for column,width in zip(table.columns,widths):column.width=width
            borders=OxmlElement('w:tblBorders')
            for edge in ('top','left','bottom','right','insideH','insideV'):
                border=OxmlElement('w:'+edge);border.set(qn('w:val'),'single');border.set(qn('w:sz'),'4');border.set(qn('w:color'),'D9D9D9');borders.append(border)
            table._tbl.tblPr.append(borders)
            for n,row in enumerate(value):
                cells=table.rows[0].cells if n==0 else table.add_row().cells
                for cell,text,width in zip(cells,row,widths):
                    cell.width=width
                    margins=OxmlElement('w:tcMar')
                    for side in ('top','left','bottom','right'):
                        margin=OxmlElement('w:'+side);margin.set(qn('w:w'),'90');margin.set(qn('w:type'),'dxa');margins.append(margin)
                    cell._tc.get_or_add_tcPr().append(margins)
                    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
                    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    cell.paragraphs[0].paragraph_format.space_after=Pt(3)
                    cell.paragraphs[0].paragraph_format.space_before=Pt(3)
                    if n==0:
                        shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'E9F2F5');cell._tc.get_or_add_tcPr().append(shading)
                    add_text(cell.paragraphs[0],text)
                    if n==0:
                        for run in cell.paragraphs[0].runs:run.bold=True
                if n==0:
                    repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
            doc.add_paragraph()
        elif kind.startswith('h'):add_text(doc.add_heading('',int(kind[1])),value)
        elif kind=='bullet':add_text(doc.add_paragraph(style='List Bullet'),value)
        elif kind=='code':
            run=doc.add_paragraph().add_run(value);run.font.name='Consolas';run.font.size=Pt(10)
        else:add_text(doc.add_paragraph(),value)
    stream=io.BytesIO();doc.save(stream);return stream.getvalue()


def export_document(title,content,path,overwrite=False):
    title,content=Drafts.validate(title,content);path=Path(path)
    suffix=path.suffix.lower()
    if suffix not in ('.docx','.html','.md'):raise ValueError('Elige .docx, .html o .md.')
    if path.exists() and not overwrite:raise FileExistsError('El archivo ya existe; elige otro nombre o confirma reemplazarlo.')
    data=word_document(title,content) if suffix=='.docx' else html_document(title,content) if suffix=='.html' else ('# '+title+'\n\n'+content+'\n').encode('utf-8')
    if not overwrite:
        # Exclusive creation prevents a concurrent exporter from replacing a file.
        with path.open('xb') as output:
            try:output.write(data);output.flush();os.fsync(output.fileno())
            except Exception:
                output.close();path.unlink(missing_ok=True);raise
    else:
        fd,tmp=tempfile.mkstemp(prefix='.aulaos-export-',suffix=suffix,dir=path.parent)
        try:
            with os.fdopen(fd,'wb') as output:output.write(data);output.flush();os.fsync(output.fileno())
            os.replace(tmp,path)
        finally:
            if os.path.exists(tmp):os.unlink(tmp)
    return path
