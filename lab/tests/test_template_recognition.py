import io
import json
import tempfile
import unittest
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from core import Store
from drafts import RevisionConflict
from templates import Templates,inspect_template,render_template,package
from template_suggestions import TemplateSuggestions
from lesson_packs import LessonPacks
from test_templates import fixture

class RecognitionTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
        self.store=Store(self.root/'db');self.gid=self.store.group('A','1');self.other=self.store.group('B','1')
        self.service=Templates(self.store);self.path=self.root/'format.docx';fixture(self.path)
    def tearDown(self):self.store.close();self.tmp.cleanup()
    def read_doc(self,doc):
        out=io.BytesIO();doc.save(out);return inspect_template(out.getvalue())
    def imported(self):
        tid=self.service.import_file(self.gid,self.path,'Formato');row=self.service.get(self.gid,tid)
        f=next(f for f in json.loads(row['schema_json'])['fields'] if f['text']=='Completar')
        return tid,f
    def test_column_hints_variable_formats_and_ids_unchanged(self):
        for n in (3,6,8):
            data=fixture(self.path,n);schema=inspect_template(data)
            fields=[f for f in schema['fields'] if f['text']=='Completar']
            self.assertEqual(len(fields),n)
            self.assertEqual([f['label_hint'] for f in fields],schema['tables'][0]['first_row'])
            self.assertEqual([f['column_start'] for f in fields],list(range(1,n+1)))
            for f in fields:self.assertEqual(f['section_hint'],'Planeación del centro educativo')
            self.assertEqual(render_template(data,{}),data)
    def test_merged_header_maps_grid_not_cell_position(self):
        doc=Document();t=doc.add_table(rows=2,cols=3)
        t.cell(0,0).merge(t.cell(0,1)).text='Proceso';t.cell(0,2).text='Evidencias'
        for c in t.rows[1].cells:c.text='Vacío'
        fields=[f for f in self.read_doc(doc)['fields'] if f['text']=='Vacío']
        self.assertEqual([f['label_hint'] for f in fields],['Proceso','Proceso','Evidencias'])
    def test_inline_and_adjacent_labels_and_no_forced_identification(self):
        doc=Document();doc.add_paragraph('Docente: Nombre');doc.add_paragraph('Sin etiqueta')
        t=doc.add_table(rows=1,cols=2);t.cell(0,0).text='Propósito:';t.cell(0,1).text=''
        fields=self.read_doc(doc)['fields']
        self.assertEqual(fields[0]['label_hint'],'Docente');self.assertEqual(fields[1]['label_hint'],'')
        self.assertEqual(fields[-1]['label_hint'],'Propósito')
    def test_empty_header_not_inferred(self):
        doc=Document();t=doc.add_table(rows=2,cols=2);t.cell(0,0).text='Uno'
        t.cell(1,0).text='Dato';t.cell(1,1).text='Otro'
        self.assertTrue(all(not f['column_hint'] for f in self.read_doc(doc)['fields']))
    def test_nested_table_uses_own_header_and_fixed_height_warning(self):
        doc=Document();outer=doc.add_table(rows=2,cols=2);outer.cell(0,0).text='A';outer.cell(0,1).text='B'
        nested=outer.cell(1,0).add_table(rows=2,cols=2);nested.cell(0,0).text='Interior';nested.cell(0,1).text='Otro'
        nested.cell(1,0).text='Contenido'
        height=OxmlElement('w:trHeight');height.set(qn('w:hRule'),'exact');height.set(qn('w:val'),'200');outer.rows[1]._tr.get_or_add_trPr().append(height)
        schema=self.read_doc(doc);f=next(f for f in schema['fields'] if f['text']=='Contenido')
        self.assertEqual(f['label_hint'],'Interior');self.assertTrue(any('altura fija' in w for w in schema['warnings']))
    def test_custom_labels_scope_conflict_reset_and_original_preserved(self):
        tid,f=self.imported();original=self.service.get(self.gid,tid)['data']
        self.assertEqual(self.service.set_label(self.gid,tid,f['id'],'Actividad de mi escuela',0),1)
        with self.assertRaises(RevisionConflict):self.service.set_label(self.gid,tid,f['id'],'Obsoleto',0)
        with self.assertRaises(ValueError):self.service.set_label(self.other,tid,f['id'],'Ajeno',1)
        with self.assertRaises(ValueError):self.service.set_label(self.gid,tid,'falso','Ajeno',1)
        row=self.service.get(self.gid,tid);field=next(x for x in json.loads(row['schema_json'])['fields'] if x['id']==f['id'])
        self.assertEqual(field['label'],'Actividad de mi escuela');self.assertTrue(field['label_custom']);self.assertEqual(row['data'],original)
        self.service.set_label(self.gid,tid,f['id'],None,1)
        field=next(x for x in json.loads(self.service.get(self.gid,tid)['schema_json'])['fields'] if x['id']==f['id'])
        self.assertEqual(field['label'],f['label_hint']);self.assertFalse(field['label_custom'])
    def test_legacy_schema_upgrade_keeps_ids_versions_and_bytes(self):
        tid,f=self.imported();did=self.service.create(self.gid,tid,'Doc');self.service.save(self.gid,did,{f['id']:'Guardado'},1)
        row=self.service.get(self.gid,tid);old=json.loads(row['schema_json']);old.pop('recognition_version')
        old['fields']=[{k:x[k] for k in ('id','location','text')} for x in old['fields']]
        with self.store.con:self.store.con.execute('UPDATE document_templates SET schema_json=? WHERE id=?',(json.dumps(old),tid))
        upgraded=self.service.get(self.gid,tid)
        self.assertEqual(upgraded['data'],row['data']);self.assertEqual(json.loads(upgraded['schema_json'])['recognition_version'],1)
        self.assertEqual(self.service.document(self.gid,did)['revision'],2)
        self.assertEqual(json.loads(self.service.document(self.gid,did)['changes']),{f['id']:'Guardado'})
        self.service.export(self.gid,did,self.root/'out.docx');self.assertIn('Guardado',[x['text'] for x in inspect_template((self.root/'out.docx').read_bytes())['fields']])
    def test_context_and_proposal_capture_labels_without_rewriting_source(self):
        tid,f=self.imported();did=self.service.create(self.gid,tid,'Doc')
        self.service.set_label(self.gid,tid,f['id'],'Vocabulario local',0)
        suggestions=TemplateSuggestions(self.store)
        class Fake:
            def respond(self,messages):return json.dumps({'changes':{},'notes':'Revisión'})
        sid=suggestions.propose(self.gid,did,'Preparar clase',1,Fake())
        self.service.set_label(self.gid,tid,f['id'],'Evaluación',1)
        context=json.loads(suggestions.get(self.gid,sid)['field_context_json'])
        self.assertEqual(next(x for x in context if x['id']==f['id'])['label'],'Vocabulario local')
        _,payload=LessonPacks(self.store).context(self.gid,did,1,'Preparar clase')
        self.assertEqual(next(x for x in payload['planning']['fields'] if x['id']==f['id'])['label'],'Evaluación')
