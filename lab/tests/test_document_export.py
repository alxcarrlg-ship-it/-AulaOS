import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile
import xml.etree.ElementTree as ET
from docx import Document
from document_export import export_document,html_document,word_document

CONTENT='''# Sesión 1
Objetivo: **practicar preguntas** en inglés.
- Listen and repeat.
- Trabaja con tu compañero.
3. Preserve this number.

| Momento | Actividad |
| --- | --- |
| Inicio | Observar y preguntar |
| Cierre | Escribir una reflexión |

## Fuentes
Adaptación del docente. [W1](https://example.org/resource?a=1&b=2)
'''

class ExportTests(unittest.TestCase):
    def setUp(self):self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
    def tearDown(self):self.tmp.cleanup()
    def test_docx_opens_with_headings_tables_and_hyperlinks(self):
        data=word_document('Guía de inglés',CONTENT);doc=Document(io.BytesIO(data))
        self.assertEqual(doc.paragraphs[0].text,'Guía de inglés')
        self.assertTrue(any(p.style.name=='Heading 1' for p in doc.paragraphs))
        self.assertEqual(doc.tables[0].cell(2,1).text,'Escribir una reflexión')
        self.assertTrue(any(p.style.name=='List Bullet' for p in doc.paragraphs))
        self.assertTrue(any('3. Preserve' in p.text for p in doc.paragraphs))
        self.assertTrue(any(r.bold for p in doc.paragraphs for r in p.runs if 'practicar' in r.text))
        with ZipFile(io.BytesIO(data)) as package:
            rels=ET.fromstring(package.read('word/_rels/document.xml.rels'))
            links=[r.attrib for r in rels if r.attrib.get('Type','').endswith('/hyperlink')]
            self.assertEqual(links[0]['Target'],'https://example.org/resource?a=1&b=2')
            self.assertEqual(links[0]['TargetMode'],'External')
            self.assertNotIn('vbaProject.bin',str(package.namelist()))
    def test_html_escapes_untrusted_content_without_remote_assets(self):
        text=html_document('<title>',CONTENT+'\n<script>alert(1)</script>\n<img src=x onerror=alert(1)>').decode()
        self.assertNotIn('<script>',text);self.assertNotIn('<img ',text)
        self.assertIn('&lt;script&gt;',text);self.assertIn('Content-Security-Policy',text)
        self.assertIn('<table>',text);self.assertIn('<strong>practicar preguntas</strong>',text)
        self.assertIn('https://example.org/resource?a=1&amp;b=2',text);self.assertIn('@page{size:letter',text)
    def test_invalid_link_not_active_in_docx_or_html(self):
        body='[bad](https://127.0.0.1/private) [file](file:///secret)'
        text=html_document('Title',body).decode();self.assertNotIn('<a href=',text)
        with ZipFile(io.BytesIO(word_document('Title',body))) as package:
            rels=ET.fromstring(package.read('word/_rels/document.xml.rels'))
            self.assertFalse(any(r.attrib.get('Type','').endswith('/hyperlink') for r in rels))
    def test_export_no_overwrite_by_default(self):
        path=self.root/'output.docx';path.write_bytes(b'original')
        with self.assertRaises(FileExistsError):export_document('Title',CONTENT,path)
        self.assertEqual(path.read_bytes(),b'original')
    def test_overwrite_failure_keeps_original(self):
        path=self.root/'output.html';path.write_bytes(b'original')
        with patch('document_export.os.replace',side_effect=OSError('busy')):
            with self.assertRaises(OSError):export_document('Title',CONTENT,path,overwrite=True)
        self.assertEqual(path.read_bytes(),b'original');self.assertEqual(list(self.root.glob('.aulaos-export-*')),[])
    def test_export_all_formats_and_preserve_spanish(self):
        for suffix in ['.docx','.html','.md']:
            with self.subTest(suffix=suffix):
                path=self.root/('output'+suffix);self.assertEqual(export_document('Guía',CONTENT,path),path)
                self.assertTrue(path.stat().st_size>100)
                if suffix!='.docx':self.assertIn('reflexión',path.read_text())
    def test_reject_extension_and_invalid_content_without_writing(self):
        for name,content in [('file.pdf',CONTENT),('file.docx','bad\x01'),('file.html','')]:
            with self.subTest(name=name),self.assertRaises(ValueError):export_document('Title',content,self.root/name)
        self.assertEqual(list(self.root.iterdir()),[])
    def test_wide_and_malformed_tables_do_not_drop_text(self):
        text='| a | b | c | d | e | f | g |\n| --- | --- | --- | --- | --- | --- | --- |\n| 1 | 2 | 3 | 4 | 5 | 6 | 7 |\n| orphan |'
        doc=Document(io.BytesIO(word_document('Title',text)));body='\n'.join(p.text for p in doc.paragraphs)
        self.assertIn('g',body);self.assertIn('7',body);self.assertIn('orphan',body)
    def test_code_remains_literal_not_a_heading_or_html(self):
        text='```\n# Literal\n<script>example</script>\n```'
        doc=Document(io.BytesIO(word_document('Title',text)))
        self.assertIn('# Literal',doc.paragraphs[1].text);self.assertEqual(doc.paragraphs[1].style.name,'Normal')
        self.assertIn('&lt;script&gt;',html_document('Title',text).decode())

if __name__=='__main__':unittest.main()
