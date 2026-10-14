import io
import json
from pathlib import Path
import tempfile
import unittest
import zipfile
from docx import Document
from docx.shared import Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import base64
from core import Store
from drafts import RevisionConflict
from templates import Templates, inspect_template, render_template, package, parse, text_of, NS


def fixture(path,columns=3):
    doc=Document();section=doc.sections[0]
    if columns>=6:section.page_width=Inches(11);section.page_height=Inches(8.5)
    section.left_margin=section.right_margin=Inches(.65)
    logo=io.BytesIO(base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAADwAAAA8CAIAAAC1nk4lAAAAXElEQVR4nO3OQQ3AIAAAMcDERGIA+fOwx5ElrYLOZ5/xN+t24AvpinRFuiJdka5IV6Qr0hXpinRFuiJdka5IV6Qr0hXpinRFuiJdka5IV6Qr0hXpinRFuiJdka68idwBchHsMmcAAAAASUVORK5CYII='))
    section.header.paragraphs[0].add_run().add_picture(logo,width=Inches(.35))
    section.footer.paragraphs[0].text='Centro educativo · Formato propio'
    p=doc.add_paragraph();p.add_run('Docente: ').bold=True;p.add_run('Nombre')
    doc.add_heading('Planeación del centro educativo',0)
    table=doc.add_table(rows=3,cols=columns);table.style='Table Grid'
    for i,cell in enumerate(table.rows[0].cells):cell.text=['Momento','Actividades','Recursos','Evaluación','Tiempo','Observaciones','Producto','Seguimiento'][i]
    for cell in table.rows[1].cells:cell.text='Completar'
    table.rows[2].cells[0].merge(table.rows[2].cells[-1]).text='Comentarios generales'
    doc.add_paragraph('Texto con enlace protegido')
    p=doc.add_paragraph();link=OxmlElement('w:hyperlink');r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='No cambiar enlace';r.append(t);link.append(r);p._p.append(link)
    doc.save(path)
    return path.read_bytes()


class TemplateTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name);self.store=Store(self.root/'db.sqlite3')
        self.gid=self.store.group('A','1');self.other=self.store.group('B','1');self.service=Templates(self.store)
        self.path=self.root/'source.docx';self.data=fixture(self.path)
    def tearDown(self):self.store.close();self.tmp.cleanup()
    def import_doc(self):return self.service.import_file(self.gid,self.path,'Formato propio','Docente A')
    def test_variable_columns_and_merges(self):
        for count in (3,6,8):
            schema=inspect_template(fixture(self.path,count));self.assertEqual(schema['tables'][0]['columns'],count)
            self.assertTrue(schema['tables'][0]['merged']);self.assertEqual(schema['images'],1)
    def test_parts_and_geometry_preserved(self):
        schema=inspect_template(self.data);field=next(f for f in schema['fields'] if f['text']=='Completar')
        output=render_template(self.data,{field['id']:'Actividad nueva\nPaso dos'})
        original=package(self.data);new=package(output)
        for name in original:
            if name!='word/document.xml':self.assertEqual(original[name],new[name],name)
        old=parse(original['word/document.xml']);updated=parse(new['word/document.xml'])
        from lxml import etree
        for xpath in ('.//w:tblPr','.//w:tblGrid','.//w:tcPr','.//w:sectPr'):
            self.assertEqual([etree.tostring(n) for n in old.findall(xpath,NS)],[etree.tostring(n) for n in updated.findall(xpath,NS)])
        self.assertIn('Actividad nueva',text_of(updated));self.assertEqual(self.path.read_bytes(),self.data)
    def test_complex_paragraphs_skipped(self):
        schema=inspect_template(self.data)
        self.assertGreaterEqual(schema['skipped'],2)
        self.assertFalse(any(f['text']=='No cambiar enlace' for f in schema['fields']))
    def test_blank_and_header_footer(self):
        schema=inspect_template(self.data);field=next(f for f in schema['fields'] if 'footer' in f['id'])
        output=render_template(self.data,{field['id']:''})
        self.assertEqual(next(f for f in inspect_template(output)['fields'] if f['id']==field['id'])['text'],'')
    def test_no_changes_exact_copy(self):self.assertEqual(render_template(self.data,{}),self.data)
    def test_invalid_keys_and_characters(self):
        field=inspect_template(self.data)['fields'][0]['id']
        for changes in ({'fake':'x'},{field:'\x01'},{field:'\ud800'},{field:42}):
            with self.assertRaises(ValueError):render_template(self.data,changes)
    def test_scope(self):
        tid=self.import_doc();did=self.service.create(self.gid,tid,'Doc')
        for fn,args in ((self.service.get,(self.other,tid)),(self.service.document,(self.other,did)),(self.service.export,(self.other,did,self.root/'x.docx'))):
            with self.assertRaises(ValueError):fn(*args)
        self.assertEqual(self.service.list(self.other),[]);self.assertEqual(self.service.documents(self.other),[])
    def test_versions_and_conflicts(self):
        tid=self.import_doc();did=self.service.create(self.gid,tid,'Doc');key=inspect_template(self.data)['fields'][0]['id']
        self.assertEqual(self.service.save(self.gid,did,{key:'Nuevo'},1),2)
        with self.assertRaises(RevisionConflict):self.service.save(self.gid,did,{key:'Viejo'},1)
        self.assertEqual(self.service.document(self.gid,did,1)['changes'],'{}')
        self.assertEqual(self.service.save(self.gid,did,{},2),3)
    def test_original_and_existing_output_never_overwritten(self):
        tid=self.import_doc();did=self.service.create(self.gid,tid,'Doc')
        with self.assertRaises(ValueError):self.service.export(self.gid,did,self.path)
        out=self.root/'copy.docx';self.service.export(self.gid,did,out)
        with self.assertRaises(FileExistsError):self.service.export(self.gid,did,out)
        self.assertEqual(out.read_bytes(),self.data)
    def test_imported_bytes_survive_source_loss_and_backup(self):
        tid=self.import_doc();did=self.service.create(self.gid,tid,'Doc');self.path.unlink()
        self.service.export(self.gid,did,self.root/'copy.docx');self.assertEqual((self.root/'copy.docx').read_bytes(),self.data)
        backup=self.store.backup();other=Store(backup)
        try:self.assertEqual(Templates(other).get(self.gid,tid)['data'],self.data)
        finally:other.close()
    def test_reject_external_resources(self):
        parts=package(self.data)
        parts['word/_rels/extra.xml.rels']=b'<Relationships><Relationship TargetMode="External" Type="image" Target="https://example.com/a.png"/></Relationships>'
        out=io.BytesIO()
        with zipfile.ZipFile(out,'w') as archive:
            for k,v in parts.items():archive.writestr(k,v)
        with self.assertRaises(ValueError):inspect_template(out.getvalue())
    def test_reject_entity_xml(self):
        out=io.BytesIO()
        with zipfile.ZipFile(out,'w') as archive:
            archive.writestr('[Content_Types].xml','<Types/>')
            archive.writestr('word/document.xml','<!DOCTYPE x [<!ENTITY a "x">]><x/>')
        with self.assertRaises(ValueError):inspect_template(out.getvalue())
    def test_split_runs_retain_first_style(self):
        schema=inspect_template(self.data);field=next(f for f in schema['fields'] if f['text']=='Docente: Nombre')
        out=render_template(self.data,{field['id']:'Docente: Otra persona'})
        doc=Document(io.BytesIO(out));self.assertEqual(doc.paragraphs[0].text,'Docente: Otra persona');self.assertTrue(doc.paragraphs[0].runs[0].bold)
    def test_independent_template_variants(self):
        one=self.import_doc();fixture(self.path,8);two=self.import_doc()
        self.assertNotEqual(self.service.get(self.gid,one)['digest'],self.service.get(self.gid,two)['digest'])
        self.assertEqual(json.loads(self.service.get(self.gid,one)['schema_json'])['tables'][0]['columns'],3)
