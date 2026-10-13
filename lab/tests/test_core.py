import tempfile
import unittest
from pathlib import Path
from core import Store

class CoreTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.db=Store(Path(self.temp.name)/'local.sqlite3')
        self.gid=self.db.group('Escuela de prueba','2 D')
        self.sid=self.db.add_student(self.gid,'Alumno Ejemplo',7)
    def tearDown(self):self.db.close();self.temp.cleanup()
    def test_mark_counts(self):
        self.db.mark(self.gid,self.sid,'2026-10-12','F')
        self.db.mark(self.gid,self.sid,'2026-10-13','R')
        self.assertEqual(self.db.attendance_counts(self.sid),{'F':1,'R':1})
        self.assertIn('1 faltas',self.db.consult(self.gid,'Cuantas faltas lleva Alumno Ejemplo'))
    def test_audit_event(self):
        self.db.log_event(self.gid,self.sid,'2026-10-12','Conducta','No completó la actividad indicada')
        self.assertIn('No completó',self.db.student_summary(self.sid))
    def test_tasks(self):
        tid=self.db.create_task('Subir planeación',self.gid,'2026-10-15')
        self.assertIn('Subir planeación',self.db.consult(self.gid,'Pendientes'))
        self.db.task_done(tid)
        self.assertEqual('No hay pendientes.',self.db.consult(self.gid,'Pendientes'))
    def test_csv_import(self):
        p=Path(self.temp.name)/'lista.csv'
        p.write_text('Nombre\nRosa Ejemplo\nMateo Ejemplo\n',encoding='utf-8')
        self.db.import_csv(self.gid,p)
        self.assertEqual(len(self.db.students(self.gid)),3)
    def test_conflict_preserves_manual(self):
        try:from openpyxl import Workbook
        except ImportError:self.skipTest('openpyxl no instalado')
        from openpyxl import Workbook
        p=Path(self.temp.name)/'lista.xlsx'
        w=Workbook();s=w.active;s.title='Asistencia'
        s.append(['N.º','Nombre del alumno(a)','12/10/2026']);s.append([1,'Rosa Ejemplo','F']);w.save(p)
        self.db.import_excel(self.gid,p,True)
        stu=self.db.find_student(self.gid,'rosa')[0]; self.db.mark(self.gid,stu['id'],'2026-10-12','A')
        self.db.import_excel(self.gid,p,True)
        self.assertEqual(self.db.marks(self.gid,'2026-10-12')[stu['id']],'A')
        self.assertEqual(len(self.db.unresolved_conflicts(self.gid)),1)

if __name__=='__main__':unittest.main()

class ExcelReorderTests(unittest.TestCase):
    def test_excel_reorder_preserves_identity(self):
        from openpyxl import Workbook,load_workbook
        with tempfile.TemporaryDirectory() as d:
            db=Store(Path(d)/'local.sqlite3');gid=db.group('Escuela X','2A')
            f=Path(d)/'lista.xlsx';w=Workbook();w.active.title='Asistencia'
            s=w.active;s.append(['N.º','Nombre del alumno(a)','12/10/2026'])
            s.append([1,'Brandon Ejemplo','A']);s.append([2,'José Ejemplo','F']);w.save(f)
            db.import_excel(gid,f,True)
            sid=db.find_student(gid,'Brandon')[0]['id']
            db.log_event(gid,sid,'2026-10-12','Conducta','Prueba de incidencia')
            s.cell(2,2).value='José Ejemplo';s.cell(2,3).value='F'
            s.cell(3,2).value='Brandon Ejemplo';s.cell(3,3).value='A';w.save(f)
            db.import_excel(gid,f,True)
            self.assertEqual(db.find_student(gid,'Brandon')[0]['id'],sid)
            self.assertIn('Prueba de incidencia',db.student_summary(sid))
            db.close()

class BackupTests(unittest.TestCase):
    def test_backup_has_data(self):
        import sqlite3
        with tempfile.TemporaryDirectory() as d:
            db=Store(Path(d)/'lab.sqlite3');db.group('Plantel','3 A')
            f=db.backup();self.assertTrue(f.exists())
            cn=sqlite3.connect(f)
            try:self.assertEqual(cn.execute('SELECT count(*) FROM groups').fetchone()[0],1)
            finally:cn.close()
            db.close()

class AmbiguousIdentityTests(unittest.TestCase):
    def test_duplicate_names_are_rejected_without_side_effects(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as d:
            db=Store(Path(d)/'lab.sqlite3'); gid=db.group('Plantel','2A')
            f=Path(d)/'lista.xlsx'; w=Workbook();w.active.title='Asistencia'
            w.active.append(['N.º','Nombre del alumno(a)','12/10/2026'])
            w.active.append([1,'Juan Ejemplo','A']);w.active.append([2,'Juan Ejemplo','F']);w.save(f)
            with self.assertRaisesRegex(ValueError,'nombres idénticos'):
                db.import_excel(gid,f,True)
            self.assertEqual(len(db.students(gid)),0)
            db.close()


class ContextAndCsvRegressionTests(unittest.TestCase):
    setUp=CoreTests.setUp
    tearDown=CoreTests.tearDown
    def test_cross_group_writes_rejected_without_changes(self):
        other=self.db.group('Otra escuela','2 D')
        self.db.mark(self.gid,self.sid,'2026-10-12','A')
        with self.assertRaises(ValueError):
            self.db.mark(other,self.sid,'2026-10-12','F')
        with self.assertRaises(ValueError):
            self.db.log_event(other,self.sid,'2026-10-12','Nota','Incorrecto')
        self.assertEqual(self.db.marks(self.gid,'2026-10-12'),{self.sid:'A'})
        self.assertEqual(self.db.marks(other,'2026-10-12'),{})
        self.assertEqual(self.db.events(self.sid),[])

    def test_csv_reorder_preserves_history_and_source(self):
        import hashlib
        p=Path(self.temp.name)/'lista.csv'
        p.write_text('Nombre\nRosa Ejemplo\nMateo Ejemplo\n',encoding='utf-8')
        self.db.import_csv(self.gid,p)
        rosa=self.db.find_student(self.gid,'rosa')[0]['id']
        self.db.mark(self.gid,rosa,'2026-10-12','F')
        self.db.log_event(self.gid,rosa,'2026-10-12','Nota','De Rosa')
        p.write_text('Nombre\nMateo Ejemplo\nRosa Ejemplo\n',encoding='utf-8')
        before=hashlib.sha256(p.read_bytes()).hexdigest()
        self.db.import_csv(self.gid,p)
        self.assertEqual(self.db.find_student(self.gid,'rosa')[0]['id'],rosa)
        mateo=self.db.find_student(self.gid,'mateo')[0]['id']
        self.assertEqual(self.db.attendance_counts(rosa),{'F':1})
        self.assertEqual(self.db.events(mateo),[])
        self.assertEqual(self.db.attendance_counts(mateo),{})
        self.assertEqual(hashlib.sha256(p.read_bytes()).hexdigest(),before)

    def test_csv_duplicate_failure_is_atomic(self):
        p=Path(self.temp.name)/'lista.csv'
        p.write_text('Nombre\nRosa Ejemplo\nRosa Ejemplo\n',encoding='utf-8')
        before=[dict(x) for x in self.db.students(self.gid)]
        with self.assertRaises(ValueError):self.db.import_csv(self.gid,p)
        self.assertEqual([dict(x) for x in self.db.students(self.gid)],before)
